from io import BytesIO

import pandas as pd
from openpyxl import load_workbook

from backend.core.workforce import generate_summary_excel, parse_workforce_xlsx


def _workbook_bytes():
    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        pd.DataFrame(
            [
                ["Jour", "LUNDI"],
                ["Date", "2026-09-21"],
                ["Type de quart", "JOUR"],
                [],
                ["Département", "Catégorie", "Cible", "Présences", "Écart"],
                ["Unité A", "PAB", 4, 2, -2],
                ["Unité B", "Inf", 3, 5, 2],
            ]
        ).to_excel(writer, sheet_name="LUNDI", header=False, index=False)
    output.seek(0)
    return output


def _workbook_with_header_on_line_three():
    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        pd.DataFrame(
            [
                ["Jour", "MARDI"],
                ["Date", "2026-09-22"],
                ["Département", "Catégorie", "Cible", "Présences", "Écart"],
                ["Unité C", "Aux", 5, 4, -1],
            ]
        ).to_excel(writer, sheet_name="MARDI", header=False, index=False)
    output.seek(0)
    return output


def test_parse_workforce_xlsx_splits_needs_and_surplus():
    needs, surplus, summary = parse_workforce_xlsx(_workbook_bytes())

    assert len(needs) == 1
    assert needs.iloc[0]["Besoins"] == 2
    assert needs.iloc[0]["Quart"] == "JOUR"
    assert len(surplus) == 1
    assert surplus.iloc[0]["Surplus"] == 2
    assert summary["lignes_analysees"] == 2


def test_parse_workforce_xlsx_uses_pandas_index_two_as_header():
    needs, surplus, summary = parse_workforce_xlsx(_workbook_with_header_on_line_three())

    assert len(needs) == 1
    assert needs.iloc[0]["Catégorie"] == "Aux"
    assert needs.iloc[0]["Besoins"] == 1
    assert surplus.empty
    assert summary["lignes_analysees"] == 1


def test_generate_summary_excel_has_expected_sheets_and_formatting():
    needs, surplus, _ = parse_workforce_xlsx(_workbook_bytes())

    workbook = load_workbook(generate_summary_excel(needs, surplus))

    assert workbook.sheetnames == ["Sommaire Besoins", "Sommaire Surplus"]
    assert workbook["Sommaire Besoins"]["A1"].value == "Jour"
    assert len(workbook["Sommaire Besoins"].conditional_formatting) == 1
    assert len(workbook["Sommaire Surplus"].conditional_formatting) == 1