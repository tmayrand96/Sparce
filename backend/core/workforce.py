"""Workforce replacement activity transformation and analysis helpers."""

from __future__ import annotations

import json
import os
import re
from io import BytesIO
from typing import Any, Dict, Optional, Tuple

import pandas as pd
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Font, PatternFill

from .summarizer import GoogleGeminiSummarizer


OUTPUT_COLUMNS = [
    "Jour",
    "Date",
    "Quart",
    "Département",
    "Catégorie",
    "Cible",
    "Présences",
    "Écart",
    "Besoins",
    "Surplus",
]


def _normalized(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip().lower())


def _find_metadata(rows: list[list[Any]], sheet_name: str) -> dict[str, Any]:
    metadata: dict[str, Any] = {"Jour": sheet_name, "Date": "", "Quart": ""}
    for row in rows[:12]:
        for index, value in enumerate(row):
            label = _normalized(value).rstrip(":")
            if label in {"jour", "date", "quart", "type de quart", "shift"}:
                next_value = row[index + 1] if index + 1 < len(row) else ""
                if label == "jour":
                    metadata["Jour"] = next_value or sheet_name
                elif label == "date":
                    metadata["Date"] = next_value
                else:
                    metadata["Quart"] = next_value
            elif label in {"jour", "date", "quart", "type de quart"}:
                metadata[label.title()] = value

    if not metadata["Quart"]:
        for candidate in (sheet_name, *(str(cell) for row in rows[:8] for cell in row)):
            match = re.search(r"\b(jour|soir|nuit)\b", _normalized(candidate))
            if match:
                metadata["Quart"] = match.group(1).upper()
                break
    return metadata


def _find_table_header(df_raw: pd.DataFrame) -> tuple[int, pd.DataFrame, dict[str, int]]:
    aliases = {
        "Département": {"département", "departement", "department", "unité", "unite"},
        "Catégorie": {"catégorie", "categorie", "catégorie d'emploi", "emploi", "poste"},
        "Cible": {"cible", "besoin", "requis", "effectif requis"},
        "Présences": {"présences", "presences", "présence", "presence", "effectif"},
        "Écart": {"écart", "ecart", "différence", "difference", "delta"},
    }
    def find_positions(columns: Any) -> dict[str, int]:
        normalized = [_normalized(cell) for cell in columns]
        positions: dict[str, int] = {}
        for name, accepted in aliases.items():
            for column_index, value in enumerate(normalized):
                if value in accepted or any(alias in value for alias in accepted if len(alias) > 5):
                    positions[name] = column_index
                    break
        return positions

    header_row_idx = None
    if len(df_raw.index) > 2:
        row_3 = df_raw.iloc[2]
        if any("département" in _normalized(value) for value in row_3):
            header_row_idx = 2

    if header_row_idx is None:
        for row_index in range(min(6, len(df_raw.index))):
            row = df_raw.iloc[row_index]
            if any("département" in _normalized(value) for value in row):
                header_row_idx = row_index
                break

    if header_row_idx is None:
        raise ValueError("Structure invalide: colonne Département introuvable dans les lignes 0 à 5.")

    df = df_raw.copy()
    df.columns = df_raw.iloc[header_row_idx].values
    df_data = df_raw.iloc[header_row_idx + 1:].copy()
    positions = find_positions(df.columns)
    if len(positions) != len(aliases):
        raise ValueError("Structure invalide: colonnes Département, Catégorie, Cible, Présences et Écart introuvables.")
    return header_row_idx, df_data, positions


def _number(value: Any) -> float:
    if pd.isna(value) or value == "":
        raise ValueError("valeur numérique manquante")
    if isinstance(value, str):
        value = value.replace(" ", "").replace(",", ".")
    return float(value)


def parse_workforce_xlsx(uploaded_file) -> Tuple[pd.DataFrame, pd.DataFrame, dict]:
    """Parse all workbook sheets into needs, surplus, and metadata summaries."""
    if uploaded_file is None:
        raise ValueError("Aucun fichier Excel fourni.")

    uploaded_file.seek(0) if hasattr(uploaded_file, "seek") else None
    try:
        workbook = pd.ExcelFile(uploaded_file)
    except Exception as exc:
        raise ValueError(f"Impossible de lire le fichier Excel: {exc}") from exc

    records: list[dict[str, Any]] = []
    sheets: list[dict[str, Any]] = []
    for sheet_name in workbook.sheet_names:
        df_raw = pd.read_excel(workbook, sheet_name=sheet_name, header=None)
        if df_raw.empty:
            continue
        rows = df_raw.values.tolist()
        metadata = _find_metadata(rows, sheet_name)
        header_index, df_data, positions = _find_table_header(df_raw)
        sheets.append(metadata.copy())
        for row in df_data.values.tolist():
            values = {name: row[index] if index < len(row) else None for name, index in positions.items()}
            if all(pd.isna(values[name]) or values[name] == "" for name in positions):
                continue
            if pd.isna(values["Département"]) or pd.isna(values["Catégorie"]):
                continue
            try:
                target = _number(values["Cible"])
                presence = _number(values["Présences"])
                difference = _number(values["Écart"])
            except (TypeError, ValueError) as exc:
                raise ValueError(f"Ligne invalide dans la feuille {sheet_name}: {exc}") from exc
            if difference == 0:
                difference = presence - target
            records.append(
                {
                    **metadata,
                    "Département": str(values["Département"]).strip(),
                    "Catégorie": str(values["Catégorie"]).strip(),
                    "Cible": target,
                    "Présences": presence,
                    "Écart": difference,
                    "Besoins": abs(difference) if difference < 0 else 0,
                    "Surplus": difference if difference > 0 else 0,
                }
            )

    if not records:
        raise ValueError("Aucune ligne d'activité exploitable dans le classeur.")
    frame = pd.DataFrame(records, columns=OUTPUT_COLUMNS)
    df_besoins = frame[frame["Besoins"] > 0].reset_index(drop=True)
    df_surplus = frame[frame["Surplus"] > 0].reset_index(drop=True)
    summary = {
        "feuilles": sheets,
        "lignes_analysees": len(frame),
        "besoins": df_besoins.to_dict(orient="records"),
        "surplus": df_surplus.to_dict(orient="records"),
    }
    return df_besoins, df_surplus, summary


def generate_summary_excel(df_besoins: pd.DataFrame, df_surplus: pd.DataFrame) -> BytesIO:
    """Create a formatted in-memory workbook containing needs and surplus."""
    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        df_besoins.to_excel(writer, sheet_name="Sommaire Besoins", index=False)
        df_surplus.to_excel(writer, sheet_name="Sommaire Surplus", index=False)
        for sheet_name, color in (("Sommaire Besoins", "FFC7CE"), ("Sommaire Surplus", "C6EFCE")):
            worksheet = writer.book[sheet_name]
            worksheet.freeze_panes = "A2"
            worksheet.auto_filter.ref = worksheet.dimensions
            for cell in worksheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="1F4E78")
            last_row = worksheet.max_row
            if last_row > 1:
                value_column = "I" if sheet_name == "Sommaire Besoins" else "J"
                worksheet.conditional_formatting.add(
                    f"{value_column}2:{value_column}{last_row}",
                    CellIsRule(operator="greaterThan", formula=["0"], fill=PatternFill("solid", fgColor=color)),
                )
            for column in worksheet.columns:
                letter = column[0].column_letter
                width = min(max(len(str(cell.value or "")) for cell in column) + 2, 32)
                worksheet.column_dimensions[letter].width = width
    output.seek(0)
    return output


def query_gemini_analysis(data_summary: dict, user_prompt: str = "", api_key: Optional[str] = None) -> str:
    """Ask Gemini for an executive workforce report based only on aggregated data."""
    if not isinstance(data_summary, dict) or not data_summary.get("lignes_analysees"):
        raise ValueError("Le contexte d'analyse est vide ou invalide.")
    default_instruction = (
        "Produis un rapport exécutif Markdown en français. Identifie les départements sous-effectifs "
        "critiques par quart, compare les ratios de couverture par catégorie (Inf, Aux, PAB, AA), "
        "et recommande des compensations entre unités en surplus et en besoin. Cite les chiffres fournis "
        "et ne fabrique aucune donnée."
    )
    instruction = user_prompt.strip() if isinstance(user_prompt, str) and user_prompt.strip() else default_instruction
    context = json.dumps(data_summary, ensure_ascii=False, default=str, indent=2)
    try:
        summarizer = GoogleGeminiSummarizer(api_key=api_key or os.getenv("GOOGLE_API_KEY"))
        return summarizer.summarize(
            f"Données agrégées de gestion des activités de remplacement:\n{context}",
            max_output_tokens=1200,
            system_instruction=instruction,
        )
    except Exception as exc:
        raise RuntimeError(f"Analyse Gemini impossible: {exc}") from exc